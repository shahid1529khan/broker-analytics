from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response, StreamingResponse
from typing import Any, Dict, Optional, List
from uuid import UUID
import io
import openpyxl
from backend.api.deps import get_current_user_with_org, verify_client_access
from backend.services.analytics_engine import AnalyticsEngine

router = APIRouter()

@router.get("/{client_id}/report.xlsx")
def export_excel(
    client_id: UUID,
    periods: Optional[List[str]] = Query(None),
    current_user: Dict[str, Any] = Depends(get_current_user_with_org)
):
    str_client_id = verify_client_access(str(client_id), current_user)
    engine = AnalyticsEngine(str_client_id, periods)
    
    if not engine.rows:
        raise HTTPException(status_code=404, detail="No data available for export")
        
    wb = openpyxl.Workbook()
    
    # Sheet 1: Dashboard Summary Data
    ws1 = wb.active
    ws1.title = "Analytics Summary"
    
    analytics = engine.get_all_analytics()
    
    ws1.append(["Metric", "Value"])
    ws1.append(["Total Book Size", f"${analytics['book_size']['trend_data'][-1]['total_balance']:,.2f}"] if analytics['book_size']['trend_data'] else ["Total Book Size", "$0"])
    ws1.append(["Growth Trend", analytics['book_size'].get('growth_trend', 'N/A')])
    ws1.append(["Average Loan Size", f"${analytics['average_loan_size']['overall_average']:,.2f}"])
    ws1.append([])
    
    # Sheet 2: Raw Normalised Rows
    ws2 = wb.create_sheet(title="Loan Data")
    headers = [
        "Period", "Aggregator", "Lender", "Loan ID", "Borrower Ref", 
        "Settlement Date", "Original Amount", "Outstanding Balance",
        "Trail Rate %", "Trail Income", "Upfront Commission"
    ]
    ws2.append(headers)
    
    for row in engine.rows:
        ws2.append([
            row.get('period_month'),
            row.get('aggregator_name'),
            row.get('lender_name'),
            row.get('loan_id'),
            row.get('borrower_reference'),
            row.get('settlement_date'),
            row.get('loan_amount_original', 0),
            row.get('outstanding_balance', 0),
            row.get('trail_rate_percent', 0),
            row.get('trail_income_this_period', 0),
            row.get('upfront_commission', 0)
        ])
        
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=yield_report_{client_id}.xlsx"}
    )

@router.get("/{client_id}/report.pdf")
def export_pdf(
    client_id: UUID,
    periods: Optional[List[str]] = Query(None),
    current_user: Dict[str, Any] = Depends(get_current_user_with_org)
):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except ImportError:
        raise HTTPException(status_code=500, detail="PDF generation not available (reportlab missing)")
        
    str_client_id = verify_client_access(str(client_id), current_user)
    engine = AnalyticsEngine(str_client_id, periods)
    
    if not engine.rows:
        raise HTTPException(status_code=404, detail="No data available for export")
        
    analytics = engine.get_all_analytics()
    
    book_size = 0
    if analytics['book_size']['trend_data']:
        book_size = analytics['book_size']['trend_data'][-1]['total_balance']
        
    stream = io.BytesIO()
    doc = SimpleDocTemplate(
        stream,
        pagesize=A4,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Yield Platform Valuation Report", styles["Title"]),
        Paragraph("Generated automatically based on normalized statement data.", styles["BodyText"]),
        Spacer(1, 14),
    ]

    def add_table(title: str, rows: List[List[Any]], col_widths: Optional[List[float]] = None):
        story.append(Paragraph(title, styles["Heading2"]))
        table = Table(rows, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(table)
        story.append(Spacer(1, 14))

    add_table("Executive Summary", [
        ["Metric", "Value"],
        ["Total Book Size", f"${book_size:,.2f}"],
        ["Growth Trend", analytics["book_size"].get("growth_trend", "N/A")],
        ["Average Loan Size", f"${analytics['average_loan_size']['overall_average']:,.2f}"],
    ], [2.2 * inch, 3.8 * inch])

    book_rows = [["Period", "Total Balance", "Net Change"]]
    for item in analytics.get("book_size", {}).get("trend_data", [])[-12:]:
        book_rows.append([
            item["period"],
            f"${item['total_balance']:,.2f}",
            f"${item.get('net_change', 0):,.2f}",
        ])
    add_table("Book Size Trend", book_rows, [1.3 * inch, 2.35 * inch, 2.35 * inch])

    trail_rows = [["Period", "Trail Income", "MoM Change"]]
    for item in analytics.get("trail_income", {}).get("trend_data", [])[-12:]:
        trail_rows.append([
            item["period"],
            f"${item['trail_income']:,.2f}",
            f"{item.get('mom_change_pct', 0)}%",
        ])
    add_table("Trail Income Trend", trail_rows, [1.3 * inch, 2.35 * inch, 2.35 * inch])

    lender_rows = [["Lender", "Balance", "% of Book"]]
    for item in analytics["lender_concentration"]["ranked_table"][:10]:
        lender_rows.append([
            item["lender_name"],
            f"${item['total_balance']:,.2f}",
            f"{item['percentage']}%",
        ])
    add_table("Lender Concentration", lender_rows, [2.6 * inch, 2.1 * inch, 1.3 * inch])

    doc.build(story)
    pdf_bytes = stream.getvalue()
    
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=yield_report_{client_id}.pdf"}
    )
